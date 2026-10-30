"""Formula and reference parsing.

The key idea is R1C1-relative normalisation: =A2*B2 in C2 and =A3*B3 in C3
are the same formula copied down, so both normalise to RC[-2]*RC[-1].
Comparing normalised forms is how the checks find a formula that breaks a
pattern, and how the diff tells a moved formula from a changed one.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from functools import lru_cache

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils.cell import column_index_from_string, get_column_letter

_CELL = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d+)$")
_COL = re.compile(r"^(\$?)([A-Za-z]{1,3})$")
_ROW = re.compile(r"^(\$?)(\d+)$")
_A1 = re.compile(r"^([A-Z]{1,3})(\d+)$")

MAX_ROW = 1_048_576
MAX_COL = 16_384

VOLATILE = {"NOW", "TODAY", "RAND", "RANDBETWEEN", "RANDARRAY", "OFFSET", "INDIRECT", "CELL", "INFO"}

# Numbers too common in formulas to be worth flagging: counts, months,
# percent, days, thousands.
BENIGN_NUMBERS = {0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 24, 60, 100, 365, 1000}


def parse_cell(ref: str) -> tuple[int, int]:
    """'C7' -> (7, 3). Raises ValueError for anything else."""
    m = _A1.match(ref.upper())
    if not m:
        raise ValueError(f"not a cell reference: {ref!r}")
    return int(m.group(2)), column_index_from_string(m.group(1))


def cell_name(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def split_sheet(operand: str) -> tuple[str, str]:
    """Split "'My Sheet'!A1:B2" into ("'My Sheet'!", "A1:B2").

    The split is at the last "!" outside quotes; quotes inside a sheet name
    are doubled, which keeps the toggle in step.
    """
    in_quote = False
    cut = -1
    for i, ch in enumerate(operand):
        if ch == "'":
            in_quote = not in_quote
        elif ch == "!" and not in_quote:
            cut = i
    if cut == -1:
        return "", operand
    return operand[: cut + 1], operand[cut + 1 :]


def unquote_sheet(prefix: str) -> str:
    """"'My Sheet'!" -> "My Sheet"; "Data!" -> "Data"."""
    name = prefix[:-1] if prefix.endswith("!") else prefix
    if len(name) >= 2 and name[0] == "'" and name[-1] == "'":
        name = name[1:-1].replace("''", "'")
    return name


def external_book(prefix: str) -> str | None:
    """The workbook part of an external reference, or None.

    "[1]Prices!" -> "1" (an index into the workbook's external links);
    "'C:\\Data\\[Prices.xlsx]Sheet1'!" -> "C:\\Data\\Prices.xlsx".
    """
    name = unquote_sheet(prefix)
    if "[" not in name or "]" not in name:
        return None
    path, rest = name.split("[", 1)
    book = rest.split("]", 1)[0]
    return path + book if path else book


def _endpoint_r1c1(part: str, row: int, col: int) -> str | None:
    m = _CELL.match(part)
    if m:
        cabs, cletters, rabs, rnum = m.groups()
        r, c = int(rnum), column_index_from_string(cletters.upper())
        return _axis("R", r, row, bool(rabs)) + _axis("C", c, col, bool(cabs))
    m = _COL.match(part)
    if m:
        cabs, cletters = m.groups()
        return _axis("C", column_index_from_string(cletters.upper()), col, bool(cabs))
    m = _ROW.match(part)
    if m:
        rabs, rnum = m.groups()
        return _axis("R", int(rnum), row, bool(rabs))
    return None


def _axis(letter: str, value: int, origin: int, absolute: bool) -> str:
    if absolute:
        return f"{letter}{value}"
    delta = value - origin
    return letter if delta == 0 else f"{letter}[{delta}]"


def ref_to_r1c1(ref: str, row: int, col: int) -> str | None:
    """'A1:B$2' seen from (row, col) -> 'R[..]C[..]:R2C[..]'; None if ref isn't A1-style."""
    parts = ref.split(":")
    if len(parts) > 2:
        return None
    out = [_endpoint_r1c1(p, row, col) for p in parts]
    if any(p is None for p in out):
        return None
    return ":".join(out)


@dataclass(frozen=True)
class _Tok:
    type: str
    subtype: str
    value: str


@lru_cache(maxsize=65536)
def _tokenize(text: str) -> tuple[_Tok, ...] | Exception:
    try:
        return tuple(_Tok(t.type, t.subtype, t.value) for t in Tokenizer(text).items)
    except Exception as exc:  # cached too, so a bad formula is parsed once
        return exc


def _tokens(formula: str) -> tuple[_Tok, ...]:
    """Tokens of a formula. Each check asks for them, so they are cached."""
    result = _tokenize(formula if formula.startswith("=") else "=" + formula)
    if isinstance(result, Exception):
        raise result
    return result


@dataclass(frozen=True)
class Ref:
    """One reference operand inside a formula."""

    text: str
    sheet: str | None  # None: the formula's own sheet
    book: str | None  # external workbook ("1" or a path), None if internal
    area: tuple[int, int, int, int] | None  # (row1, col1, row2, col2); None for names etc.


def _col(letters: str) -> int:
    return column_index_from_string(letters.upper())


def _area(ref: str) -> tuple[int, int, int, int] | None:
    """(row1, col1, row2, col2) covered by an A1 reference; None for names etc."""
    parts = ref.split(":")
    if len(parts) == 1:
        m = _CELL.match(parts[0])
        if not m:
            return None
        r, c = int(m.group(4)), _col(m.group(2))
        return (r, c, r, c)
    if len(parts) != 2:
        return None
    a, b = parts
    ma, mb = _CELL.match(a), _CELL.match(b)
    if ma and mb:
        r1, c1, r2, c2 = int(ma.group(4)), _col(ma.group(2)), int(mb.group(4)), _col(mb.group(2))
        return (min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))
    ma, mb = _COL.match(a), _COL.match(b)
    if ma and mb:  # whole columns, e.g. A:C
        c1, c2 = _col(ma.group(2)), _col(mb.group(2))
        return (1, min(c1, c2), MAX_ROW, max(c1, c2))
    ma, mb = _ROW.match(a), _ROW.match(b)
    if ma and mb:  # whole rows, e.g. 3:5
        r1, r2 = int(ma.group(2)), int(mb.group(2))
        return (min(r1, r2), 1, max(r1, r2), MAX_COL)
    return None


def references(formula: str) -> list[Ref]:
    """All reference operands in a formula (cells, ranges, names, structured refs)."""
    refs = []
    for tok in _tokens(formula):
        if tok.type == Token.OPERAND and tok.subtype == Token.RANGE:
            prefix, ref = split_sheet(tok.value)
            book = external_book(prefix) if prefix else None
            sheet = unquote_sheet(prefix) if prefix and book is None else None
            refs.append(Ref(tok.value, sheet, book, _area(ref)))
    return refs


def _canonical_prefix(prefix: str, sheet_alias: dict[str, str] | None) -> tuple[str, str | None]:
    """(prefix spelled one way, sheet name or None for external refs)."""
    if not prefix or external_book(prefix) is not None:
        return prefix.upper(), None
    name = unquote_sheet(prefix)
    for old, new in (sheet_alias or {}).items():
        if name.lower() == old.lower():
            name = new
            break
    return ("'" + name.replace("'", "''") + "'!").upper(), name


def _rewrite(formula: str, ref_fn) -> str:
    """Rebuild a formula with ref_fn(prefix, ref) applied to each reference.

    Whitespace is dropped and everything except text literals is
    upper-cased, since Excel ignores both. Unparseable formulas come back
    upper-cased.
    """
    try:
        tokens = _tokens(formula)
    except Exception:
        return formula.upper()
    out = []
    for tok in tokens:
        if tok.type == Token.WSPACE:
            continue
        if tok.type == Token.OPERAND and tok.subtype == Token.RANGE:
            prefix, ref = split_sheet(tok.value)
            out.append(ref_fn(prefix, ref))
        elif tok.type == Token.OPERAND and tok.subtype == Token.TEXT:
            out.append(tok.value)
        else:
            out.append(tok.value.upper())
    return "".join(out)


def normalize(formula: str, row: int, col: int, sheet_alias: dict[str, str] | None = None,
              own_sheet: str | None = None) -> str:
    """R1C1-relative form of a formula, for comparing copies of it.

    Every reference becomes relative to (row, col), the way a formula looks
    when it is copied down or across. With own_sheet set, only references
    to that sheet become relative and references to other sheets stay as
    written: the form for spotting a formula that moved because rows were
    inserted, since inserting rows doesn't change references to other
    sheets. sheet_alias renames sheets, so a formula that points at a
    renamed sheet still matches.
    """

    def fn(prefix: str, ref: str) -> str:
        canon, name = _canonical_prefix(prefix, sheet_alias)
        if own_sheet is not None and prefix and (name is None or name.lower() != own_sheet.lower()):
            return canon + ref.upper()
        r1c1 = ref_to_r1c1(ref, row, col)
        return canon + (r1c1 if r1c1 is not None else ref.upper())

    return _rewrite(formula, fn)


def _map_endpoint(part: str, axis: int, line_map: dict[int, int]) -> str:
    m = _CELL.match(part)
    if m:
        cabs, letters, rabs, num = m.groups()
        r, c = int(num), _col(letters)
        if axis == 0:
            r = line_map.get(r, r)
        else:
            c = line_map.get(c, c)
        return f"{cabs}{get_column_letter(c)}{rabs}{r}"
    m = _COL.match(part)
    if m and axis == 1:
        cabs, letters = m.groups()
        c = _col(letters)
        return f"{cabs}{get_column_letter(line_map.get(c, c))}"
    m = _ROW.match(part)
    if m and axis == 0:
        rabs, num = m.groups()
        return f"{rabs}{line_map.get(int(num), int(num))}"
    return part.upper()


def remap(formula: str, own_sheet: str, maps: dict[str, tuple[str, int, dict[int, int]]],
          sheet_alias: dict[str, str] | None = None) -> str:
    """A formula from the old version, with each reference moved to where
    its cells are in the new version.

    maps: old sheet name (lower case) -> (new sheet name, axis, line map),
    where axis 0 maps rows and axis 1 maps columns. Compare the result with
    canonical() of the new formula.
    """

    def fn(prefix: str, ref: str) -> str:
        canon, name = _canonical_prefix(prefix, sheet_alias)
        if prefix and name is None:  # another workbook
            return canon + ref.upper()
        old_name = unquote_sheet(prefix) if prefix else own_sheet
        target = maps.get(old_name.lower())
        if target is None:
            return canon + ref.upper()
        _, axis, line_map = target
        return canon + ":".join(_map_endpoint(p, axis, line_map) for p in ref.split(":"))

    return _rewrite(formula, fn)


def shape(formula: str, sheet_alias: dict[str, str] | None = None) -> str:
    """A formula with its cell references blanked out: =A2*B2 and =C5*D9
    both become ?*?. Sheet names, functions and constants are kept. Used to
    line up rows and columns, which must survive inserted rows and columns."""
    return _rewrite(formula, lambda prefix, ref: _canonical_prefix(prefix, sheet_alias)[0] + ("?" if _area(ref) else ref.upper()))


def canonical(formula: str) -> str:
    """A formula spelled one way (no whitespace, upper case, quoted sheet names)."""
    return _rewrite(formula, lambda prefix, ref: _canonical_prefix(prefix, None)[0] + ref.upper())


def functions(formula: str) -> list[str]:
    """Function names used, without Excel's _xlfn./_xlws. storage prefixes."""
    names = []
    try:
        tokens = _tokens(formula)
    except Exception:
        return names
    for tok in tokens:
        if tok.type == Token.FUNC and tok.subtype == Token.OPEN:
            name = tok.value[:-1].upper()
            for p in ("_XLFN.", "_XLWS."):
                if name.startswith(p):
                    name = name[len(p) :]
            names.append(name)
    return names


def numbers(formula: str) -> list[str]:
    """Number literals typed into a formula (as written)."""
    try:
        tokens = _tokens(formula)
    except Exception:
        return []
    return [t.value for t in tokens if t.type == Token.OPERAND and t.subtype == Token.NUMBER]


def has_ref_error(formula: str) -> bool:
    """True if #REF! appears in the formula outside text literals."""
    try:
        tokens = _tokens(formula)
    except Exception:
        return "#REF!" in formula.upper()
    return any(t.subtype != Token.TEXT and "#REF!" in t.value.upper() for t in tokens)


def is_benign_number(literal: str) -> bool:
    try:
        value = float(literal)
    except ValueError:
        return True
    return value.is_integer() and abs(int(value)) in BENIGN_NUMBERS
