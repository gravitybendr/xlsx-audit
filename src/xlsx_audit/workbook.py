"""Load a workbook into a small model: formulas, typed values and cached results.

openpyxl reads the cells (twice: once for formulas, once for the results
Excel saved). Sheet visibility, defined names and external links are read
from the package XML directly, which keeps them independent of openpyxl
internals.
"""
from __future__ import annotations

import posixpath
import warnings
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .refs import cell_name

MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
REL = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class WorkbookError(Exception):
    """The file can't be audited (missing, encrypted, wrong format)."""


@dataclass
class Cell:
    row: int
    col: int
    formula: str | None = None  # without the leading "="
    value: object = None  # typed value, for cells without a formula
    cached: object = None  # last result Excel saved for a formula
    error: str | None = None  # error the cell shows (#REF!, #N/A...), if any
    array: bool = False

    @property
    def name(self) -> str:
        return cell_name(self.row, self.col)

    @property
    def is_number(self) -> bool:
        import datetime

        v = self.value
        return (
            self.formula is None
            and self.error is None
            and not isinstance(v, bool)
            and isinstance(v, (int, float, datetime.date, datetime.time, datetime.timedelta))
        )


@dataclass
class Sheet:
    name: str
    state: str = "visible"  # visible, hidden, veryHidden
    kind: str = "worksheet"  # worksheet or chartsheet
    cells: dict[tuple[int, int], Cell] = field(default_factory=dict)

    @property
    def max_row(self) -> int:
        return max((r for r, _ in self.cells), default=0)

    @property
    def max_col(self) -> int:
        return max((c for _, c in self.cells), default=0)


@dataclass
class DefinedName:
    name: str
    text: str
    scope: str | None = None  # sheet name for sheet-level names


@dataclass
class Workbook:
    path: Path
    sheets: list[Sheet]
    names: list[DefinedName]
    external_links: list[str]  # target of [1], [2]... in formulas

    def sheet(self, name: str) -> Sheet | None:
        for s in self.sheets:
            if s.name.lower() == name.lower():
                return s
        return None

    @property
    def stats(self) -> dict:
        cells = sum(len(s.cells) for s in self.sheets)
        formulas = sum(1 for s in self.sheets for c in s.cells.values() if c.formula is not None)
        return {"sheets": len(self.sheets), "cells": cells, "formulas": formulas}


def _check_file(path: Path) -> None:
    if not path.exists():
        raise WorkbookError(f"{path}: file not found")
    if path.is_dir():
        raise WorkbookError(f"{path}: is a folder, not a workbook")
    with path.open("rb") as fh:
        head = fh.read(8)
    if head == OLE_MAGIC:
        raise WorkbookError(
            f"{path}: this is a password-protected workbook or an old .xls file. "
            "Remove the password, or save it as .xlsx in Excel, and try again."
        )
    if not zipfile.is_zipfile(path):
        raise WorkbookError(f"{path}: not an .xlsx or .xlsm file")
    if path.suffix.lower() == ".xlsb":
        raise WorkbookError(f"{path}: binary .xlsb workbooks aren't supported; save as .xlsx in Excel")


def _rels(z: zipfile.ZipFile, part: str) -> dict[str, tuple[str, str, str]]:
    """Relationships of a package part: rId -> (type, target, mode)."""
    folder, name = posixpath.split(part)
    rels_path = posixpath.join(folder, "_rels", name + ".rels")
    if rels_path not in z.namelist():
        return {}
    root = ET.fromstring(z.read(rels_path))
    out = {}
    for rel in root.iter(f"{PKG}Relationship"):
        target = rel.get("Target", "")
        mode = rel.get("TargetMode", "Internal")
        if mode != "External":
            target = posixpath.normpath(posixpath.join(folder, target)) if not target.startswith("/") else target.lstrip("/")
        out[rel.get("Id")] = (rel.get("Type", ""), target, mode)
    return out


def _read_package(path: Path) -> tuple[list[tuple[str, str, str]], list[DefinedName], list[str]]:
    """Sheets (name, state, kind), defined names and external link targets."""
    with zipfile.ZipFile(path) as z:
        root_rels = _rels(z, "")
        wb_part = next((t for typ, t, _ in root_rels.values() if typ.endswith("/officeDocument")), "xl/workbook.xml")
        root = ET.fromstring(z.read(wb_part))
        rels = _rels(z, wb_part)

        sheets = []
        for el in root.iter(f"{MAIN}sheet"):
            typ = rels.get(el.get(f"{REL}id"), ("", "", ""))[0]
            kind = "chartsheet" if typ.endswith("/chartsheet") else "worksheet"
            sheets.append((el.get("name", ""), el.get("state", "visible"), kind))

        names = []
        for el in root.iter(f"{MAIN}definedName"):
            scope = None
            if el.get("localSheetId") is not None:
                idx = int(el.get("localSheetId"))
                scope = sheets[idx][0] if idx < len(sheets) else None
            names.append(DefinedName(el.get("name", ""), (el.text or "").strip(), scope))

        links = []
        for el in root.iter(f"{MAIN}externalReference"):
            _typ, part, _mode = rels.get(el.get(f"{REL}id"), ("", "", ""))
            links.append(_external_target(z, part))
    return sheets, names, links


def _external_target(z: zipfile.ZipFile, part: str) -> str:
    if not part or part not in z.namelist():
        return "(unknown)"
    root = ET.fromstring(z.read(part))
    book = root.find(f"{MAIN}externalBook")
    if book is not None:
        rel = _rels(z, part).get(book.get(f"{REL}id"))
        return rel[1] if rel else "(unknown)"
    dde = root.find(f"{MAIN}ddeLink")
    if dde is not None:
        return f"DDE link {dde.get('ddeService', '')}|{dde.get('ddeTopic', '')}"
    ole = root.find(f"{MAIN}oleLink")
    if ole is not None:
        rel = _rels(z, part).get(ole.get(f"{REL}id"))
        return f"OLE link {rel[1] if rel else ''}".strip()
    return "(unknown)"


def _formula_text(value) -> tuple[str | None, bool]:
    """(formula without "=", is_array) for an openpyxl cell value."""
    if isinstance(value, str) and value.startswith("=") and len(value) > 1:
        return value[1:], False
    text = getattr(value, "text", None)
    if isinstance(text, str):  # ArrayFormula
        return text[1:] if text.startswith("=") else text, True
    if type(value).__name__ == "DataTableFormula":
        return "TABLE()", True
    return None, False


def load(path: str | Path) -> Workbook:
    """Read a workbook. Raises WorkbookError for files that can't be audited."""
    path = Path(path)
    _check_file(path)
    try:
        sheet_info, names, links = _read_package(path)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")  # openpyxl warns about unsupported extensions
            wb_f = openpyxl.load_workbook(path, data_only=False, keep_links=False)
            wb_v = openpyxl.load_workbook(path, data_only=True, keep_links=False)
    except WorkbookError:
        raise
    except (zipfile.BadZipFile, KeyError, ET.ParseError, OSError, ValueError, TypeError) as exc:
        raise WorkbookError(f"{path}: can't read workbook ({exc})") from exc

    states = {name: (state, kind) for name, state, kind in sheet_info}
    sheets = []
    for ws in wb_f.worksheets:
        state, kind = states.get(ws.title, (ws.sheet_state, "worksheet"))
        sheet = Sheet(ws.title, state, kind)
        values = wb_v[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                formula, is_array = _formula_text(c.value)
                cell = Cell(c.row, c.column, formula=formula, array=is_array)
                if formula is None:
                    cell.value = c.value
                    if c.data_type == "e":
                        cell.error = str(c.value)
                else:
                    v = values.cell(row=c.row, column=c.column)
                    cell.cached = v.value
                    if v.data_type == "e" and v.value is not None:
                        cell.error = str(v.value)
                sheet.cells[(c.row, c.column)] = cell
        sheets.append(sheet)
    # Chartsheets have no cells but still count for visibility checks.
    for name, state, kind in sheet_info:
        if kind == "chartsheet" and not any(s.name == name for s in sheets):
            sheets.append(Sheet(name, state, kind))
    order = {name: i for i, (name, _, _) in enumerate(sheet_info)}
    sheets.sort(key=lambda s: order.get(s.name, len(order)))
    return Workbook(path, sheets, names, links)
